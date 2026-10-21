import docx

doc = docx.Document('Project 4 TECR_Net_revised.docx')

def get_math_text(element):
    math_texts = []
    for node in element.iter():
        if node.tag.endswith('}t'): # m:t or w:t
            math_texts.append(node.text or '')
    return ''.join(math_texts)

with open('scratch/math_extracted.txt', 'w', encoding='utf-8') as f:
    for i in range(len(doc.paragraphs)):
        p = doc.paragraphs[i]
        xml_str = p._element.xml
        if '<m:oMath' in xml_str:
            m_text = get_math_text(p._element)
            f.write(f"Para {i}: text='{p.text}' | MathContent='{m_text}'\n")

print("Done writing scratch/math_extracted.txt")
