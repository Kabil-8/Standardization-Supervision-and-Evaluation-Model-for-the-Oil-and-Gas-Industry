import docx
import re

doc = docx.Document('Project 4 TECR_Net_revised.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")

eq_indices = []
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    m = re.match(r'^\((\d+)\)$', txt)
    if m:
        eq_indices.append((i, txt, [r.text for r in p.runs]))
        # check if there are images or xml elements inside the paragraph
        xml_str = p._element.xml
        has_drawing = '<w:drawing' in xml_str or '<w:pict' in xml_str
        has_omath = '<m:oMath' in xml_str
        print(f"Eq para {i}: text='{txt}', runs={len(p.runs)}, drawing={has_drawing}, oMath={has_omath}")

print(f"Found {len(eq_indices)} bare equation numbers.")
